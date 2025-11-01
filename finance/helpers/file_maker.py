from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from io import BytesIO
from django.http import HttpResponse, JsonResponse, FileResponse

def _generate_excel_response(data_dict, filename="export.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "팀 재무데이터"

    # ✅ 월 헤더
    month_labels = [f"{i}월" for i in range(1, 13)]
    month_keys = ["jan", "feb", "mar", "apr", "may", "jun",
                  "jul", "aug", "sep", "oct", "nov", "dec"]

    # ✅ 첫 줄 (상단 헤더: 매출/매출총이익)
    header_row1 = (
        [""] * 12  # 앞쪽 기본 정보 칼럼은 빈 문자열
        + ["매출"] * 12
        + ["매출총이익"] * 12
    )

    # ✅ 두 번째 줄 (1월~12월 x 2)
    header_row2 = (
        ["구분", "DIV", "그룹", "프로젝트 분류", "운영 팀", "협업 팀", "Job Code", "프로젝트명",
         "광고주 코드", "광고주명", "수익 유형", "기획협업 본부"]
        + month_labels  # 매출
        + month_labels  # 매출총이익
    )

    # ✅ 시트에 헤더 추가
    ws.append(header_row1)
    ws.append(header_row2)

    # ✅ 헤더 스타일 지정
    for cell in ws["1:1"]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in ws["2:2"]:
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # ✅ 데이터 행 구성
    def flatten_row(obj, is_tp=True):
        md = obj.monthly_data.model_dump()
        return [
            obj.category,
            obj.div,
            obj.group_name,
            obj.project_category,
            obj.operation_team.team_name if obj.operation_team else "",
            obj.shared_team.team_name if is_tp and getattr(obj, "shared_team", None) else "",
            obj.project_code,
            obj.project_detail,
            obj.ship_to_party,
            obj.client_name,
            obj.revenue_type,
            obj.headquarter,
            *[md.get(m, {}).get("sales", 0) for m in month_keys],
            *[md.get(m, {}).get("gross_profit", 0) for m in month_keys],
        ]

    for obj in data_dict["team_predictions"]:
        ws.append(flatten_row(obj, is_tp=True))
    for obj in data_dict["sap_processed"]:
        ws.append(flatten_row(obj, is_tp=False))

    # ✅ 셀 너비 자동 조정
    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max(10, max_length + 2)

    # ✅ 저장 및 응답
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return FileResponse(
        output,
        as_attachment=True,
        filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

def _make_name(division: str = None, group: str = None, team: str = None) -> str:
    if not division:
        return "all"
    if not group:
        return division
    if not team:
        return f"{division}-{group}"
    return f"{division}-{group}-{team}"
